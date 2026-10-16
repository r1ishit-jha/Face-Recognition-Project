"""RISHIT JHA | KU2407U181 | 202407020172 | CSE-D"""
import argparse
import copy
import csv
import hashlib
import json
import re
import shutil
import time
from pathlib import Path

import numpy as np
import torch
from torch import nn

ROOT = Path(__file__).resolve().parent
UNKNOWN = 'unidentified'


def normalize(x):
    return x / np.maximum(np.linalg.norm(x, axis=-1, keepdims=True), 1e-12)


def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2), encoding='utf-8')


def engine(device):
    from facenet_pytorch import MTCNN, InceptionResnetV1
    detector = MTCNN(image_size=160, margin=20, keep_all=True,
                     post_process=False, device=device)
    encoder = InceptionResnetV1(pretrained='vggface2').eval().to(device)
    encoder.requires_grad_(False)
    return detector, encoder


def extract(image, detector, encoder, device):
    from facenet_pytorch.models.mtcnn import fixed_image_standardization
    boxes, probabilities = detector.detect(image)
    if boxes is None:
        return np.empty((0, 4)), np.empty((0, 512), dtype=np.float32)
    good = np.asarray(probabilities, dtype=float) >= 0.95
    boxes = boxes[good]
    if len(boxes) == 0:
        return boxes, np.empty((0, 512), dtype=np.float32)
    faces = detector.extract(image, boxes, save_path=None)
    faces = fixed_image_standardization(faces).to(device)
    with torch.inference_mode():
        features = encoder(faces).cpu().numpy()
    return boxes, normalize(features)


def capture(args):
    import cv2
    from PIL import Image
    from facenet_pytorch import MTCNN
    name = args.name.strip()
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9 _-]{0,59}', name):
        raise ValueError('Use a simple name with letters, spaces, digits, _ or -.')
    if name == UNKNOWN and args.split == 'train':
        raise ValueError('Unknown people belong only in validate or test.')
    folder = args.data / args.split / name
    folder.mkdir(parents=True, exist_ok=True)
    detector = MTCNN(keep_all=True, device=args.device)
    cap = cv2.VideoCapture(args.camera)
    if not cap.isOpened():
        raise RuntimeError('Camera unavailable. Close other camera apps or change --camera.')
    count, last = 0, 0.0
    session = time.time_ns()
    try:
        while count < args.count:
            ok, frame = cap.read()
            if not ok:
                raise RuntimeError('Camera frame could not be read.')
            frame = cv2.resize(frame, (640, 480))
            raw = frame.copy()
            boxes, probs = detector.detect(Image.fromarray(cv2.cvtColor(raw, cv2.COLOR_BGR2RGB)))
            valid = boxes is not None and len(boxes) == 1 and probs[0] >= 0.95
            if valid and time.monotonic() - last >= args.interval:
                path = folder / f'{session}_{count:04d}.jpg'
                if not cv2.imwrite(str(path), raw):
                    raise OSError(f'Could not save {path}')
                count += 1
                last = time.monotonic()
            cv2.putText(frame, f'{name}: {count}/{args.count} - Q stops', (10, 30),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.65, (0, 255, 0), 2)
            cv2.putText(frame, 'One clear face only. Change pose and lighting.', (10, 60),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 255, 255), 1)
            cv2.imshow('Collect face dataset', frame)
            if cv2.waitKey(1) & 0xff == ord('q'):
                break
    finally:
        cap.release()
        cv2.destroyAllWindows()
    print(f'Saved {count} images to {folder}')


def dataset(args, split, detector, encoder):
    from PIL import Image
    features, names, paths, rejected = [], [], [], []
    for path in sorted((args.data / split).glob('*/*')):
        if path.suffix.lower() not in {'.jpg', '.jpeg', '.png'}:
            continue
        try:
            with Image.open(path) as image:
                boxes, embeddings = extract(image.convert('RGB'), detector, encoder, args.device)
            if len(boxes) != 1:
                rejected.append({'path': str(path), 'reason': 'Expected exactly one confident face'})
                continue
            features.append(embeddings[0])
            names.append(path.parent.name)
            paths.append(str(path.relative_to(args.data)))
        except (OSError, ValueError) as error:
            rejected.append({'path': str(path), 'reason': str(error)})
    save_json(args.output / f'{split}_rejected.json', rejected)
    if not features:
        raise ValueError(f'No usable images in {args.data / split}')
    return np.asarray(features, dtype=np.float32), np.asarray(names), paths


def audit(data):
    seen, counts = {}, {}
    for split in ('train', 'validate', 'test'):
        counts[split] = {}
        for p in sorted((data / split).glob('*/*')):
            if p.suffix.lower() not in {'.jpg', '.jpeg', '.png'}:
                continue
            digest = hashlib.sha256(p.read_bytes()).hexdigest()
            if digest in seen and seen[digest][0] != split:
                raise ValueError(f'Exact duplicate across splits: {p} and {seen[digest][1]}')
            seen[digest] = (split, str(p))
            counts[split][p.parent.name] = counts[split].get(p.parent.name, 0) + 1
    return counts


def fit_head(x, y, vx, vy, epochs=150):
    torch.manual_seed(42)
    head = nn.Linear(x.shape[1], int(y.max()) + 1)
    opt = torch.optim.Adam(head.parameters(), lr=0.01, weight_decay=0.001)
    loss_fn = nn.CrossEntropyLoss()
    x, vx = torch.tensor(x), torch.tensor(vx)
    y, vy = torch.tensor(y, dtype=torch.long), torch.tensor(vy, dtype=torch.long)
    history, best, state, stale = [], float('inf'), None, 0
    for epoch in range(epochs):
        head.train()
        for batch in torch.randperm(len(x)).split(32):
            opt.zero_grad()
            loss_fn(head(x[batch]), y[batch]).backward()
            opt.step()
        head.eval()
        with torch.no_grad():
            a, b = head(x), head(vx)
            tl, vl = loss_fn(a, y).item(), loss_fn(b, vy).item()
            ta, va = (a.argmax(1) == y).float().mean().item(), (b.argmax(1) == vy).float().mean().item()
        history.append({'epoch': epoch + 1, 'train_loss': tl, 'val_loss': vl,
                        'train_accuracy': ta, 'val_accuracy': va})
        if vl < best - 1e-5:
            best, state, stale = vl, copy.deepcopy(head.state_dict()), 0
        else:
            stale += 1
        if stale >= 15:
            break
    head.load_state_dict(state)
    return head.eval(), history


def scores(x, weights, bias, centers):
    logits = x @ weights.T + bias
    logits -= logits.max(axis=1, keepdims=True)
    probability = np.exp(logits)
    probability /= probability.sum(axis=1, keepdims=True)
    ids = probability.argmax(axis=1)
    distance = np.linalg.norm(x - centers[ids], axis=1)
    return ids, probability.max(axis=1), distance


def decide(ids, probability, distance, classes, pmin, dmax):
    return np.where((probability >= pmin) & (distance <= dmax),
                    np.asarray(classes)[ids], UNKNOWN)


def calibrate(x, labels, weights, bias, centers, classes):
    ids, prob, dist = scores(x, weights, bias, centers)
    known = labels != UNKNOWN
    if not known.any() or known.all():
        raise ValueError('Validation needs known members and consented unknown people.')
    best = None
    for d in np.linspace(0.3, 1.6, 66):
        for p in np.linspace(0.25, 0.95, 29):
            pred = decide(ids, prob, dist, classes, p, d)
            objective = ((pred[known] == labels[known]).mean() +
                         (pred[~known] == UNKNOWN).mean()) / 2
            candidate = (float(objective), -float(d), float(p))
            if best is None or candidate > best[0]:
                best = (candidate, float(p), float(d))
    return {'validation_balanced_accept_reject': best[0][0],
            'probability_min': best[1], 'distance_max': best[2]}


def fit_single(x, vx, labels, name):
    center = normalize(x.mean(axis=0, keepdims=True))
    distance = np.linalg.norm(vx - center[0], axis=1)
    known = labels == name
    if not known.any() or known.all():
        raise ValueError('Validation needs your own photos and unidentified photos of another person.')
    values = np.unique(distance)
    candidates = np.unique(np.r_[0.0, (values[:-1] + values[1:]) / 2, 2.0])
    best = max(((float(((distance[known] <= d).mean() +
                       (distance[~known] > d).mean()) / 2), -float(d))
                for d in candidates))
    config = {'mode': 'single_person_centroid', 'probability_min': None,
              'distance_max': -best[1], 'validation_balanced_accept_reject': best[0]}
    arrays = (np.zeros((1, x.shape[1]), dtype=np.float32),
              np.zeros(1, dtype=np.float32), center)
    return config, arrays


def train(args):
    counts = audit(args.data)
    print('Reading face images. Please wait...', flush=True)
    detector, encoder = engine(args.device)
    x, names, _ = dataset(args, 'train', detector, encoder)
    vx, vnames, _ = dataset(args, 'validate', detector, encoder)
    classes = sorted(set(names))
    if UNKNOWN in classes:
        raise ValueError('Keep unidentified photos in validate/test only, not data/train.')
    if set(vnames) != set(classes) | {UNKNOWN}:
        missing = (set(classes) | {UNKNOWN}) - set(vnames)
        extra = set(vnames) - (set(classes) | {UNKNOWN})
        raise ValueError(f'Validation labels do not match. Missing: {sorted(missing)}; extra: {sorted(extra)}. Check rejected-image logs too.')
    if any((names == c).sum() < 5 for c in classes):
        raise ValueError('Collect at least five usable training images per member; 60 recommended.')
    known = vnames != UNKNOWN
    history = []
    if len(classes) == 1:
        print(f'Single-person mode: enrolling {classes[0]} and calibrating rejection.', flush=True)
        config, (weights, bias, centers) = fit_single(x, vx, vnames, classes[0])
    else:
        y = np.array([classes.index(n) for n in names])
        vy = np.array([classes.index(n) for n in vnames[known]])
        head, history = fit_head(x, y, vx[known], vy, args.epochs)
        weights = head.weight.detach().numpy()
        bias = head.bias.detach().numpy()
        centers = normalize(np.stack([x[names == c].mean(0) for c in classes]))
        config = calibrate(vx, vnames, weights, bias, centers, classes)
        config['mode'] = 'multi_person_classifier'
    config.update({'classes': classes, 'encoder': 'InceptionResnetV1-vggface2',
                   'seed': 42, 'epochs_run': len(history),
                   'best_epoch': min(history, key=lambda h: h['val_loss'])['epoch'] if history else None,
                   'raw_counts': counts, 'usable_train': len(x), 'usable_validation': len(vx),
                   'model_id': str(time.time_ns())})
    args.output.mkdir(parents=True, exist_ok=True)
    np.savez(args.output / 'classifier.npz', weights=weights, bias=bias, centers=centers)
    save_json(args.output / 'model.json', config)
    save_json(args.output / 'history.json', history)
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    if history:
        fig, axes = plt.subplots(1, 2, figsize=(10, 4))
        for ax, metric in zip(axes, ('loss', 'accuracy')):
            for split in ('train', 'val'):
                ax.plot([h['epoch'] for h in history], [h[f'{split}_{metric}'] for h in history], label=split)
            ax.set(xlabel='Epoch', ylabel=metric.title())
            ax.legend()
        graph = 'training_curves.png'
    else:
        distance = np.linalg.norm(vx - centers[0], axis=1)
        fig, ax = plt.subplots(figsize=(9, 4))
        ax.scatter(np.flatnonzero(known), distance[known], label=classes[0], marker='o')
        ax.scatter(np.flatnonzero(~known), distance[~known], label='Unknown validation person', marker='x')
        ax.axhline(config['distance_max'], color='red', linestyle='--', label='Selected rejection threshold')
        ax.set(xlabel='Validation image index', ylabel='Distance to enrolled face centroid',
               title='Single-person calibration - no classifier epochs')
        ax.legend()
        graph = 'validation_distances.png'
    fig.tight_layout()
    fig.savefig(args.output / graph, dpi=160)
    plt.close(fig)
    print(json.dumps(config, indent=2))
    print('Model saved. You can now run the live command.')
    if len(classes) == 1:
        print('This mode enrolls a face template; it does not train a one-class softmax network.')
        print('Your unknown validation photos calibrate rejection, but do not prove rejection of every unfamiliar person.')


def load_model(output):
    if not (output / 'model.json').exists():
        raise ValueError('No saved model yet. Run train successfully before evaluate or live.')
    config = json.loads((output / 'model.json').read_text())
    with np.load(output / 'classifier.npz', allow_pickle=False) as z:
        arrays = tuple(z[k].copy() for k in ('weights', 'bias', 'centers'))
    return config, arrays


def predict(x, config, arrays):
    if config.get('mode') == 'single_person_centroid':
        distance = np.linalg.norm(x - arrays[2][0], axis=1)
        labels = np.where(distance <= config['distance_max'], config['classes'][0], UNKNOWN)
        return labels, np.full(len(x), np.nan), distance
    ids, probability, distance = scores(x, *arrays)
    labels = decide(ids, probability, distance, config['classes'],
                    config['probability_min'], config['distance_max'])
    return labels, probability, distance


def evaluate(args):
    from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, ConfusionMatrixDisplay
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    audit(args.data)
    config, arrays = load_model(args.output)
    detector, encoder = engine(args.device)
    start = time.perf_counter()
    x, actual, paths = dataset(args, 'test', detector, encoder)
    elapsed = time.perf_counter() - start
    if set(actual) != set(config['classes']) | {UNKNOWN}:
        raise ValueError('Test must include all members and unidentified, with matching folder names.')
    predicted, prob, dist = predict(x, config, arrays)
    labels = config['classes'] + [UNKNOWN]
    report = classification_report(actual, predicted, labels=labels, output_dict=True, zero_division=0)
    known = actual != UNKNOWN
    report.update({'model_id': config.get('model_id'), 'mode': config.get('mode', 'multi_person_classifier'),
                   'open_set_accuracy': float(accuracy_score(actual, predicted)),
                   'known_identification_accuracy': float((actual[known] == predicted[known]).mean()),
                   'unknown_false_accept_rate': float((predicted[~known] != UNKNOWN).mean()),
                   'known_rejection_rate': float((predicted[known] == UNKNOWN).mean()),
                   'accepted_test_images': len(actual),
                   'dataset_processing_ms_per_image': 1000 * elapsed / len(actual)})
    save_json(args.output / 'metrics.json', report)
    with (args.output / 'predictions.csv').open('w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['image', 'actual', 'predicted', 'max_softmax', 'centroid_distance'])
        writer.writerows(zip(paths, actual, predicted, ['' if np.isnan(p) else p for p in prob], dist))
    matrix = confusion_matrix(actual, predicted, labels=labels)
    np.savetxt(args.output / 'confusion_matrix.csv', matrix, fmt='%d', delimiter=',')
    fig, ax = plt.subplots(figsize=(8, 7))
    ConfusionMatrixDisplay(matrix, display_labels=labels).plot(ax=ax, xticks_rotation=35, colorbar=False)
    fig.tight_layout()
    fig.savefig(args.output / 'confusion_matrix.png', dpi=160)
    plt.close(fig)
    fig, axes = plt.subplots(2, 3, figsize=(12, 7))
    from PIL import Image
    for ax, i in zip(axes.flat, list(range(min(6, len(paths)))) + [None] * 6):
        ax.axis('off')
        if i is not None:
            with Image.open(args.data / paths[i]) as im:
                ax.imshow(im)
            ax.set_title(f'True: {actual[i]}\nPredicted: {predicted[i]}', fontsize=9)
    fig.tight_layout()
    fig.savefig(args.output / 'sample_predictions.png', dpi=140)
    plt.close(fig)
    print(json.dumps(report, indent=2))


def live(args):
    import cv2
    from PIL import Image
    config, arrays = load_model(args.output)
    detector, encoder = engine(args.device)
    cap = cv2.VideoCapture(args.camera)
    if not cap.isOpened():
        raise RuntimeError('Camera unavailable.')
    try:
        while True:
            start = time.perf_counter()
            ok, frame = cap.read()
            if not ok:
                raise RuntimeError('Camera frame could not be read.')
            frame = cv2.resize(frame, (640, 480))
            boxes, x = extract(Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)), detector, encoder, args.device)
            if len(x):
                labels, _, _ = predict(x, config, arrays)
                for box, name in zip(boxes, labels):
                    a, b, c, d = box.astype(int)
                    color = (0, 0, 255) if name == UNKNOWN else (0, 200, 0)
                    cv2.rectangle(frame, (a, b), (c, d), color, 2)
                    cv2.putText(frame, name, (max(a, 0), max(b - 8, 20)), cv2.FONT_HERSHEY_SIMPLEX, 0.65, color, 2)
            fps = 1 / max(time.perf_counter() - start, 1e-6)
            cv2.putText(frame, f'{fps:.1f} processing FPS | Q quit | S screenshot', (10, 25),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 0), 1)
            cv2.imshow('FaceNet group recognition', frame)
            key = cv2.waitKey(1) & 0xff
            if key == ord('s'):
                cv2.imwrite(str(args.output / f'live_{time.time_ns()}.png'), frame)
            if key == ord('q'):
                break
    finally:
        cap.release()
        cv2.destroyAllWindows()


def results_report(args):
    from docx import Document
    from docx.shared import Inches
    config = json.loads((args.output / 'model.json').read_text())
    metrics = json.loads((args.output / 'metrics.json').read_text())
    if config.get('model_id') != metrics.get('model_id'):
        raise ValueError('Results are from an earlier model. Run evaluate again before results-report.')
    doc = Document()
    doc.add_heading('Measured Face Recognition Results', 0)
    doc.add_paragraph('RISHIT JHA | KU2407U181 | 202407020172 | CSE-D')
    doc.add_paragraph('Generated from local training and held-out evaluation outputs. Append to the project report after checking labels and collection sessions.')
    doc.add_paragraph('Recognition mode: ' + config.get('mode', 'multi_person_classifier'))
    if config.get('mode') == 'single_person_centroid':
        doc.add_paragraph('Single-person mode enrolls a normalized centroid and calibrates a distance threshold. No classifier backpropagation or softmax probabilities are used. Zero epochs is expected.')
    for key in ('epochs_run', 'best_epoch', 'usable_train', 'usable_validation', 'probability_min', 'distance_max'):
        doc.add_paragraph(f'{key}: {config[key]}')
    for key in ('open_set_accuracy', 'known_identification_accuracy', 'unknown_false_accept_rate', 'known_rejection_rate', 'accepted_test_images', 'dataset_processing_ms_per_image'):
        doc.add_paragraph(f'{key}: {metrics[key]}')
    training_graph = 'validation_distances' if config.get('mode') == 'single_person_centroid' else 'training_curves'
    for name in (training_graph, 'confusion_matrix', 'sample_predictions'):
        doc.add_heading(name.replace('_', ' ').title(), 1)
        doc.add_picture(str(args.output / f'{name}.png'), width=Inches(6))
    doc.add_heading('Individual test cases', 1)
    rows = list(csv.DictReader((args.output / 'predictions.csv').open(encoding='utf-8')))
    for row in rows[:10]:
        doc.add_paragraph(f"{row['image']} | Actual: {row['actual']} | Predicted: {row['predicted']}")
    doc.save(args.output / 'Measured_Results.docx')
    print(args.output / 'Measured_Results.docx')


def import_images(args):
    from PIL import Image
    name = args.name.strip()
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9 _-]{0,59}', name):
        raise ValueError('Use a simple member name with letters, digits, spaces, _ or -.')
    if name == UNKNOWN and args.split == 'train':
        raise ValueError('Unknown people belong only in validate or test.')
    if args.source is None or not args.source.is_dir():
        raise ValueError('Use --source with a folder of JPG or PNG photos.')
    target = args.data / args.split / name
    if args.source.resolve() == target.resolve():
        raise ValueError('Source photos are already in this destination folder.')
    candidates = sorted(p for p in args.source.iterdir() if p.suffix.lower() in {'.jpg', '.jpeg', '.png'})
    target.mkdir(parents=True, exist_ok=True)
    existing = {}
    for split in ('train', 'validate', 'test'):
        for path in (args.data / split).glob('*/*'):
            if path.suffix.lower() in {'.jpg', '.jpeg', '.png'}:
                existing[hashlib.sha256(path.read_bytes()).hexdigest()] = path
    saved, skipped = 0, 0
    session = time.time_ns()
    for photo in candidates:
        try:
            with Image.open(photo) as im:
                im.verify()
            digest = hashlib.sha256(photo.read_bytes()).hexdigest()
            if digest in existing:
                print(f'Skipped duplicate: {photo.name} (already at {existing[digest]})')
                skipped += 1
                continue
            dest = target / f'import_{session}_{saved:04d}{photo.suffix.lower()}'
            shutil.copyfile(photo, dest)
            existing[digest] = dest
            saved += 1
        except (OSError, ValueError) as error:
            print(f'Skipped {photo.name}: {error}')
            skipped += 1
    print(f'Imported {saved} photos to {target}; skipped {skipped}.')
    print('Face detection is checked during training. Each photo should show one clear face.')
    print('Use different photos/sessions for train, validate and test, then rerun train.')


def selftest(args):
    torch.set_num_threads(1)
    rng = np.random.default_rng(42)
    basis = np.eye(8, dtype=np.float32)[:3]
    x = normalize(np.repeat(basis, 20, axis=0) + rng.normal(0, .03, (60, 8))).astype(np.float32)
    y = np.repeat(np.arange(3), 20)
    vx = normalize(np.repeat(basis, 5, axis=0) + rng.normal(0, .03, (15, 8))).astype(np.float32)
    vy = np.repeat(np.arange(3), 5)
    head, history = fit_head(x, y, vx, vy, 60)
    arrays = (head.weight.detach().numpy(), head.bias.detach().numpy(), basis)
    classes = ['A', 'B', 'C']
    ux = np.repeat(np.eye(8, dtype=np.float32)[3:4], 5, axis=0)
    validation = np.vstack([vx, ux])
    labels = np.array([classes[i] for i in vy] + [UNKNOWN] * 5)
    config = calibrate(validation, labels, *arrays, classes)
    config['classes'] = classes
    pred, _, _ = predict(validation, config, arrays)
    checks = {'head_learns_separable_features': history[-1]['train_accuracy'] == 1.0,
              'known_labels_correct': bool(np.all(pred[:15] == labels[:15])),
              'unknown_vectors_rejected': bool(np.all(pred[15:] == UNKNOWN)),
              'unit_normalization': bool(np.allclose(np.linalg.norm(x, axis=1), 1)),
              'history_loss_decreases': history[-1]['train_loss'] < history[0]['train_loss'],
              'rejection_overrides_softmax': decide(np.array([0]), np.array([.99]), np.array([2.0]), classes, .5, 1.0)[0] == UNKNOWN}
    sx = x[y == 0]
    sv = np.vstack([vx[vy == 0], ux])
    sl = np.array(['A'] * 5 + [UNKNOWN] * 5)
    single_config, single_arrays = fit_single(sx, sv, sl, 'A')
    single_config['classes'] = ['A']
    single_pred, single_prob, _ = predict(sv, single_config, single_arrays)
    checks['single_person_known_and_unknown'] = bool(np.array_equal(single_pred, sl))
    checks['single_person_no_fake_probability'] = bool(np.isnan(single_prob).all())
    checks = {k: bool(v) for k, v in checks.items()}
    save_json(args.output / 'software_checks.json', {'scope': 'Synthetic vectors only; not face accuracy or webcam validation', 'checks': checks})
    print(json.dumps(checks, indent=2))
    if not all(checks.values()):
        raise AssertionError('A software check failed.')


def main():
    parser = argparse.ArgumentParser(description='FaceNet transfer-learning mini project')
    parser.add_argument('command', choices=['capture', 'train', 'evaluate', 'live', 'results-report', 'selftest', 'import-images'])
    parser.add_argument('--source', type=Path, help='Photo folder for import-images')
    parser.add_argument('--data', type=Path, default=ROOT / 'data')
    parser.add_argument('--output', type=Path, default=ROOT / 'outputs')
    parser.add_argument('--device', default='cuda' if torch.cuda.is_available() else 'cpu')
    parser.add_argument('--name', default='Rishit Jha')
    parser.add_argument('--split', choices=['train', 'validate', 'test'], default='train')
    parser.add_argument('--camera', type=int, default=0)
    parser.add_argument('--count', type=int, default=60)
    parser.add_argument('--interval', type=float, default=0.7)
    parser.add_argument('--epochs', type=int, default=150)
    args = parser.parse_args()
    if args.count < 1 or args.epochs < 1 or args.interval < .1:
        parser.error('Use positive count/epochs and interval >= 0.1 seconds.')
    try:
        globals()[{'results-report': 'results_report', 'import-images': 'import_images', 'selftest': 'selftest'}.get(args.command, args.command)](args)
    except (ValueError, OSError, RuntimeError, ImportError) as error:
        parser.exit(1, f'Error: {error}\n')


if __name__ == '__main__':
    main()
